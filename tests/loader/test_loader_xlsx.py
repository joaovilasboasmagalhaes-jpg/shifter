import pytest
from openpyxl import Workbook

from src.loader.loader_xlsx import _read_axis_values, load_file


def test_read_axis_values_reads_row_and_skips_header():
    workbook = Workbook()
    sheet = workbook.active
    if not sheet:
        raise ValueError("Failed to create or access the active worksheet in the workbook.")

    sheet["A1"] = "Header"
    sheet["B1"] = "D1"
    sheet["C1"] = "D2"

    values = _read_axis_values(sheet, {"row": True, "index": 1, "header": True})

    assert values == ["D1", "D2"]


def test_read_axis_values_reads_column_without_header():
    workbook = Workbook()
    sheet = workbook.active
    if not sheet:
        raise ValueError("Failed to create or access the active worksheet in the workbook.")

    sheet["B1"] = "W1"
    sheet["B2"] = "W2"

    values = _read_axis_values(sheet, {"row": False, "index": 2, "header": False})

    assert values == ["W1", "W2"]


def test_read_axis_values_raises_for_invalid_index():
    workbook = Workbook()
    sheet = workbook.active
    if not sheet:
        raise ValueError("Failed to create or access the active worksheet in the workbook.")

    with pytest.raises(ValueError):
        _read_axis_values(sheet, {"row": True, "index": 0, "header": False})


def test_load_file_reads_row_dates_and_column_workers(tmp_path):
    workbook = Workbook()
    sheet = workbook.active

    if not sheet:
        raise ValueError("Failed to create or access the active worksheet in the workbook.")

    sheet["A1"] = "Header"
    sheet["B1"] = "2026-03-01"
    sheet["C1"] = "2026-03-02"
    sheet["D1"] = "2026-03-03"

    sheet["A2"] = "Alice"
    sheet["A3"] = "Bob"

    file_path = tmp_path / "row_dates.xlsx"
    workbook.save(file_path)

    config = {
        "date": {"row": True, "index": 1, "header": True},
        "worker": {"row": False, "index": 1, "header": True},
    }

    result = load_file(str(file_path), config)

    assert result["dates"] == ["2026-03-01", "2026-03-02", "2026-03-03"]
    assert result["workers"] == ["Alice", "Bob"]


def test_load_file_reads_column_dates_and_row_workers(tmp_path):
    workbook = Workbook()
    sheet = workbook.active

    if not sheet:
        raise ValueError("Failed to create or access the active worksheet in the workbook.")

    sheet["A1"] = "Header"
    sheet["B1"] = "Alice"
    sheet["C1"] = "Bob"
    sheet["D1"] = "Carol"

    sheet["A2"] = "2026-04-01"
    sheet["A3"] = "2026-04-02"

    file_path = tmp_path / "column_dates.xlsx"
    workbook.save(file_path)

    config = {
        "date": {"row": False, "index": 1, "header": True},
        "worker": {"row": True, "index": 1, "header": True},
    }

    result = load_file(str(file_path), config)

    assert result["dates"] == ["2026-04-01", "2026-04-02"]
    assert result["workers"] == ["Alice", "Bob", "Carol"]


def test_load_file_raises_for_missing_config_keys(tmp_path):
    workbook = Workbook()
    file_path = tmp_path / "missing_keys.xlsx"
    workbook.save(file_path)

    with pytest.raises(ValueError):
        load_file(str(file_path), {"date": {"row": True, "index": 1, "header": False}})
