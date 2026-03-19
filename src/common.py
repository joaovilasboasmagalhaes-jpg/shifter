import json
from contextlib import contextmanager
from datetime import date
from pathlib import Path
from typing import Any, Dict, Iterator, Optional, Tuple


def resolve_file_path(
    file_path: str,
    file_type: Optional[str] = None,
) -> Tuple[str, Optional[str]]:
    """Return a normalized file path and resolved extension.

    The extension is chosen in this order:
    1) explicit ``file_type`` argument,
    2) extension already present in ``file_path``.
    """
    if not file_path:
        raise ValueError("file_path is required")

    path = Path(file_path).expanduser()
    normalized_file_type = (file_type or "").strip().lower().lstrip(".")
    path_extension = path.suffix.lower().lstrip(".")

    resolved_extension = normalized_file_type or path_extension or None

    if not path.suffix and resolved_extension:
        path = path.with_suffix(f".{resolved_extension}")

    return str(path), resolved_extension


def validate_existing_file_path(file_path: str) -> str:
    """Validate that a file path exists and points to a file."""
    if not file_path:
        raise ValueError("file_path is required")

    path = Path(file_path).expanduser()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    return str(path)


def load_json_config(
    config_path: str, required_root_keys: Optional[Tuple[str, ...]] = None
) -> Dict[str, Any]:
    """Load a JSON config file and validate required root keys.

    If ``required_root_keys`` is provided, each key must exist in the root object.
    """
    path = Path(validate_existing_file_path(config_path))
    if path.suffix.lower() != ".json":
        raise ValueError(f"Config file must be a .json file: {path}")

    try:
        with path.open("r", encoding="utf-8") as config_file:
            config_data = json.load(config_file)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid JSON in config file {path}: {exc.msg}") from exc

    if not isinstance(config_data, dict):
        raise ValueError("Config JSON must contain an object at the root")

    for key in required_root_keys or ():
        if key not in config_data:
            raise ValueError(f"Config JSON must contain a '{key}' object")

    return config_data


@contextmanager
def open_excel_workbook(
    file_path: str,
    *,
    data_only: bool = True,
    read_only: bool = True,
) -> Iterator[Any]:
    """Open an Excel workbook and ensure it is always closed."""
    from openpyxl import load_workbook

    validated_path = validate_existing_file_path(file_path)
    workbook = load_workbook(validated_path, data_only=data_only, read_only=read_only)
    try:
        yield workbook
    finally:
        workbook.close()


def read_excel_matrix(
    sheet: Any,
    *,
    start_row: int,
    start_col: int,
    row_count: int,
    col_count: int,
) -> list[list[Any]]:
    """Read a rectangular region from a worksheet into a matrix."""
    if start_row < 1 or start_col < 1:
        raise ValueError("start_row and start_col must be greater than or equal to 1")
    if row_count < 0 or col_count < 0:
        raise ValueError("row_count and col_count must be greater than or equal to 0")

    if row_count == 0 or col_count == 0:
        return []

    end_col = start_col + col_count - 1
    matrix: list[list[Any]] = []

    for row_index in range(start_row, start_row + row_count):
        row_values = next(
            sheet.iter_rows(
                min_row=row_index,
                max_row=row_index,
                min_col=start_col,
                max_col=end_col,
                values_only=True,
            ),
            tuple(),
        )
        matrix.append(list(row_values))

    return matrix


def build_import_date(year: int, month: int, day: int) -> date:
    """Build a date from imported year/month/day values with validation."""
    if not isinstance(year, int):
        raise ValueError("Import config 'year' must be an integer")
    if not isinstance(month, int):
        raise ValueError("Import config 'month' must be an integer")
    if month < 1 or month > 12:
        raise ValueError("Import config 'month' must be between 1 and 12")
    if not isinstance(day, int):
        raise ValueError("Day value must be an integer")

    try:
        return date(year, month, day)
    except ValueError as exc:
        raise ValueError(f"Invalid day '{day}' for {year:04d}-{month:02d}") from exc
